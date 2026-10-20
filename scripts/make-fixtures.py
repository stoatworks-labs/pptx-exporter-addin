#!/usr/bin/env python3
"""Build the test decks in test/fixtures/ with python-pptx.

    npm run fixtures          (makes .venv with python-pptx + Pillow if needed)

sample-16x9.pptx is built to exercise every path the exporter reads out of the
OOXML: title and body placeholders, nested bullets with bold/italic/link runs,
a two-column layout, a table with a merged cell, a picture with alt text, a
hidden slide, transitions (plain and the p14 AlternateContent form), rehearsed
timings, a slide-jump link, LaTeX's special characters, sections, notes in
several scripts, and an embedded video. sample-4x3.pptx is two slides at 4:3,
for the sizing code.

The decks are committed; rerun this only when a test needs something new.
"""
import io
import os
import subprocess
import sys
import tempfile

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'test', 'fixtures')

NS = {
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
    'p14': 'http://schemas.microsoft.com/office/powerpoint/2010/main',
}


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def set_transition(slide, xml):
    """Insert a <p:transition> (or an mc:AlternateContent wrapping one) where
    the schema wants it: after p:clrMapOvr, before p:timing."""
    el = etree.fromstring(xml)
    sld = slide._element
    timing = sld.find('p:timing', NS)
    if timing is not None:
        timing.addprevious(el)
    else:
        ext = sld.find('p:extLst', NS)
        if ext is not None:
            ext.addprevious(el)
        else:
            sld.append(el)


def picture_png():
    img = Image.new('RGB', (640, 360), (31, 58, 96))
    d = ImageDraw.Draw(img)
    for i in range(0, 640, 40):
        d.line([(i, 0), (i + 200, 360)], fill=(240, 180, 60), width=6)
    d.ellipse([220, 80, 420, 280], outline=(255, 255, 255), width=10)
    buf = io.BytesIO()
    img.save(buf, 'PNG')
    buf.seek(0)
    return buf


def tiny_video(path):
    subprocess.run([
        'ffmpeg', '-v', 'error', '-y', '-f', 'lavfi', '-i', 'testsrc2=size=320x180:rate=25',
        '-t', '2', '-pix_fmt', 'yuv420p', '-c:v', 'libx264', '-movflags', '+faststart', path,
    ], check=True)


def build_16x9():
    prs = Presentation()
    widen_to_16x9(prs)
    prs.core_properties.title = 'Advanced Exporter Test Deck'
    prs.core_properties.author = 'Fixture Author'
    L = prs.slide_layouts  # 0 title, 1 title+content, 3 two content, 5 title only, 6 blank

    # 1 — title slide
    s = prs.slides.add_slide(L[0])
    s.shapes.title.text = 'Advanced Exporter Test Deck'
    s.placeholders[1].text = 'Fixture — généré & 日本語'
    notes(s, 'Welcome notes with ünïcödé ✓\nSecond line of the opening notes.')

    # 2 — bullets with levels and run formatting
    s = prs.slides.add_slide(L[1])
    s.shapes.title.text = 'Agenda'
    tf = s.placeholders[1].text_frame
    tf.text = 'First point'
    p = tf.add_paragraph(); p.level = 1
    r = p.add_run(); r.text = 'bold'; r.font.bold = True
    r = p.add_run(); r.text = ' and '
    r = p.add_run(); r.text = 'italic'; r.font.italic = True
    p = tf.add_paragraph(); p.level = 2; p.text = 'Third level'
    p = tf.add_paragraph(); p.level = 0
    r = p.add_run(); r.text = 'A link to example.com'
    r.hyperlink.address = 'https://example.com/'
    notes(s, 'Agenda notes.\n\nA paragraph after a blank line.')

    # 3 — two columns
    s = prs.slides.add_slide(L[3])
    s.shapes.title.text = 'Two Columns'
    s.placeholders[1].text = 'Left column'
    s.placeholders[2].text = 'Right column'

    # 4 — table with a merged cell
    s = prs.slides.add_slide(L[5])
    s.shapes.title.text = 'A Table'
    tbl = s.shapes.add_table(3, 3, Inches(1), Inches(2), Inches(8), Inches(3)).table
    for ri, row in enumerate([['Kit', 'Qty', 'Notes'], ['Projector', '2', 'Laser'], ['Screen', '1', '']]):
        for ci, val in enumerate(row):
            tbl.cell(ri, ci).text = val
    tbl.cell(2, 1).merge(tbl.cell(2, 2))
    tbl.cell(2, 1).text = 'Merged cell'
    notes(s, 'Notes in Ελληνικά and Русский.')

    # 5 — picture with alt text
    s = prs.slides.add_slide(L[5])
    s.shapes.title.text = 'A Picture'
    pic = s.shapes.add_picture(picture_png(), Inches(3), Inches(1.8), Inches(7), Inches(3.9375))
    pic._element.nvPicPr.cNvPr.set('descr', 'Diagonal stripes and a circle')

    # 6 — hidden slide
    s = prs.slides.add_slide(L[1])
    s.shapes.title.text = 'Hidden Slide'
    s.placeholders[1].text = 'You should not see this unless hidden slides are included.'
    s._element.set('show', '0')
    notes(s, 'Notes on the hidden slide.')

    # 7 — plain fade transition, auto-advance after 3 s
    s = prs.slides.add_slide(L[1])
    s.shapes.title.text = 'Fade and Advance'
    s.placeholders[1].text = 'This slide fades in and advances after three seconds.'
    set_transition(s, '<p:transition xmlns:p="%s" spd="med" advTm="3000"><p:fade/></p:transition>' % NS['p'])

    # 8 — p14 transition + slide-jump link + LaTeX specials
    s = prs.slides.add_slide(L[1])
    s.shapes.title.text = 'Specials: 50% & $5 #1'
    tf = s.placeholders[1].text_frame
    tf.text = r'Braces {x} tilde ~ caret ^ underscore _ backslash \ end'
    p = tf.add_paragraph()
    r = p.add_run(); r.text = 'Back to the agenda'
    rid = s.part.relate_to(prs.slides[1].part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide')
    rPr = r._r.get_or_add_rPr()
    h = etree.SubElement(rPr, '{%s}hlinkClick' % NS['a'])
    h.set('{%s}id' % NS['r'], rid)
    h.set('action', 'ppaction://hlinksldjump')
    set_transition(s, (
        '<mc:AlternateContent xmlns:mc="{mc}" xmlns:p14="{p14}">'
        '<mc:Choice Requires="p14"><p:transition xmlns:p="{p}" spd="slow" p14:dur="1600"><p:push dir="u"/></p:transition></mc:Choice>'
        '<mc:Fallback><p:transition xmlns:p="{p}" spd="slow"><p:push dir="u"/></p:transition></mc:Fallback>'
        '</mc:AlternateContent>').format(**NS))
    notes(s, 'Notes with specials: 100% {sure} & $cash$ #tag ~ ^ _ \\ done.')

    # 9 — embedded video
    s = prs.slides.add_slide(L[5])
    s.shapes.title.text = 'A Video'
    with tempfile.TemporaryDirectory() as td:
        mp4 = os.path.join(td, 'clip.mp4')
        tiny_video(mp4)
        poster = os.path.join(td, 'poster.png')
        Image.new('RGB', (320, 180), (0, 0, 0)).save(poster)
        s.shapes.add_movie(mp4, Inches(3.5), Inches(2), Inches(6.4), Inches(3.6), poster_frame_image=poster, mime_type='video/mp4')
    notes(s, '日本語のノート。')

    add_sections(prs, [('Intro', [0, 1]), ('Body', [2, 3, 4, 5]), ('Finale', [6, 7, 8])])
    return prs


def widen_to_16x9(prs):
    """python-pptx's template is 10 x 7.5 in. Make it 13.333 x 7.5 and stretch
    the master's and layouts' placeholders across, so PowerPoint renders the
    deck the way a real widescreen deck looks rather than hugging the left."""
    factor = 12192000 / prs.slide_width
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    sz = prs.part._element.find('p:sldSz', NS)
    if 'type' in sz.attrib:
        del sz.attrib['type']
    parts = [prs.slide_master] + list(prs.slide_layouts)
    for part in parts:
        for off in part._element.iter('{%s}off' % NS['a']):
            off.set('x', str(int(int(off.get('x')) * factor)))
        for ext in part._element.iter('{%s}ext' % NS['a']):
            if ext.get('cx') is not None:
                ext.set('cx', str(int(int(ext.get('cx')) * factor)))


def add_sections(prs, sections):
    pres = prs.part._element
    ids = [sld.get('id') for sld in pres.find('p:sldIdLst', NS)]
    ext_lst = pres.find('p:extLst', NS)
    if ext_lst is None:
        ext_lst = etree.SubElement(pres, '{%s}extLst' % NS['p'])
    ext = etree.SubElement(ext_lst, '{%s}ext' % NS['p'])
    ext.set('uri', '{521415D9-36F7-43E2-AB2F-B90AF26B5E84}')
    lst = etree.SubElement(ext, '{%s}sectionLst' % NS['p14'], nsmap={'p14': NS['p14']})
    for n, (name, idxs) in enumerate(sections):
        sec = etree.SubElement(lst, '{%s}section' % NS['p14'])
        sec.set('name', name)
        sec.set('id', '{00000000-0000-0000-0000-00000000000%d}' % n)
        sl = etree.SubElement(sec, '{%s}sldIdLst' % NS['p14'])
        for i in idxs:
            etree.SubElement(sl, '{%s}sldId' % NS['p14']).set('id', ids[i])


def build_4x3():
    prs = Presentation()  # python-pptx's default is 10 x 7.5 in, i.e. 4:3
    prs.core_properties.title = 'Four by Three'
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = 'Four by Three'
    s.placeholders[1].text = 'An older shape of screen'
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = 'Second'
    s.placeholders[1].text = 'Bullet'
    tb = s.shapes.add_textbox(Inches(1), Inches(6), Inches(4), Inches(0.6))
    tb.text_frame.text = 'A loose text box'
    tb.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xC0, 0x30, 0x30)
    notes(s, 'Four by three notes.')
    return prs


def main():
    os.makedirs(OUT, exist_ok=True)
    build_16x9().save(os.path.join(OUT, 'sample-16x9.pptx'))
    build_4x3().save(os.path.join(OUT, 'sample-4x3.pptx'))
    print('wrote', OUT)


if __name__ == '__main__':
    sys.exit(main())
